from io import BytesIO

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from docx import Document as DocxDocument

from accounts.models import OrganizationMembership, User
from candidates.bulk_intake import upload_candidate_intake_cv
from candidates.documents import DOCX_CONTENT_TYPE, upload_candidate_cv
from candidates.models import (
    Candidate,
    CandidateIntakeBatch,
    CandidateIntakeItem,
    CandidateSource,
    CandidateVacancyConsideration,
)
from candidates.reuse import add_candidates_from_pool
from organizations.models import Organization, OrganizationRetentionPolicy
from outreach.workflow import assess_contact_permission
from tests.vacancy_candidate_helpers import associate_candidate_with_vacancy
from vacancies.models import Vacancy, VacancyRequirements
from vacancies.services import confirm_requirements_and_open_vacancy

pytestmark = pytest.mark.django_db


def cv_upload(*, name="candidate.docx", email="candidate@example.test"):
    document = DocxDocument()
    document.add_paragraph("Candidate Person")
    document.add_paragraph(f"{email} | +383 44 123 456 | Prishtina")
    document.add_paragraph("Python and Django experience")
    output = BytesIO()
    document.save(output)
    return SimpleUploadedFile(name, output.getvalue(), content_type=DOCX_CONTENT_TYPE)


def workspace_with_open_vacancy():
    user = User.objects.create_user(username="vacancy-recruiter")
    organization = Organization.objects.create(name="Northstar", slug="northstar")
    OrganizationMembership.objects.create(
        user=user,
        organization=organization,
        role=OrganizationMembership.Role.ADMIN,
    )
    vacancy = Vacancy.objects.create(
        organization=organization,
        title="Senior Django Developer",
        description="Build secure Django services.",
        created_by=user,
    )
    requirements = VacancyRequirements.objects.create(
        vacancy=vacancy,
        source_description=vacancy.description,
        summary="Senior backend role",
        created_by=user,
    )
    confirm_requirements_and_open_vacancy(requirements=requirements, user=user)
    vacancy.refresh_from_db()
    return user, organization, vacancy


def intake_url(organization, vacancy):
    return reverse(
        "candidates:vacancy-candidate-intake-create",
        args=[organization.slug, vacancy.pk],
    )


def test_vacancy_page_leads_to_minimal_cv_upload_and_flags_missing_privacy_default(
    client,
):
    user, organization, vacancy = workspace_with_open_vacancy()
    client.force_login(user)

    vacancy_page = client.get(
        reverse("vacancies:vacancy-detail", args=[organization.slug, vacancy.pk])
    )
    response = client.get(intake_url(organization, vacancy))
    content = response.content.decode()

    assert response.status_code == 200
    assert intake_url(organization, vacancy) in vacancy_page.content.decode()
    assert "Add candidate CVs" in content
    assert "Shared details" not in content
    assert "Source name" not in content
    assert "Consent" not in content
    assert "This vacancy only" in content
    assert "Reason for storing applicants is not configured" in content


def test_vacancy_candidate_views_are_scoped_and_pool_remains_separate(client):
    user, organization, vacancy = workspace_with_open_vacancy()
    associated = Candidate.objects.create(
        organization=organization,
        full_name="Vacancy Candidate",
        created_by=user,
    )
    Candidate.objects.create(
        organization=organization,
        full_name="Pool Candidate",
        created_by=user,
    )
    associate_candidate_with_vacancy(
        candidate=associated,
        vacancy=vacancy,
        user=user,
    )
    client.force_login(user)

    vacancy_list = client.get(
        reverse(
            "candidates:vacancy-candidate-list",
            args=[organization.slug, vacancy.pk],
        )
    )
    vacancy_detail = client.get(
        reverse("vacancies:vacancy-detail", args=[organization.slug, vacancy.pk])
    )
    pool = client.get(reverse("candidates:candidate-list", args=[organization.slug]))

    assert vacancy_list.status_code == 200
    assert "Vacancy Candidate" in vacancy_list.content.decode()
    assert "Pool Candidate" not in vacancy_list.content.decode()
    assert "Vacancy Candidate" in vacancy_detail.content.decode()
    assert "Pool Candidate" not in vacancy_detail.content.decode()
    assert "Organization candidate pool" in vacancy_list.content.decode()
    assert "Vacancy Candidate" in pool.content.decode()
    assert "Pool Candidate" in pool.content.decode()


def test_add_from_pool_shows_permission_status_and_records_deliberate_reuse(client):
    user, organization, vacancy = workspace_with_open_vacancy()
    reusable = Candidate.objects.create(
        organization=organization,
        full_name="Reusable Candidate",
        location="Prishtina",
        created_by=user,
    )
    reusable_source = CandidateSource.objects.create(
        candidate=reusable,
        source_type=CandidateSource.SourceType.REFERRAL,
        source_name="Synthetic referral",
        lawful_basis=CandidateSource.LawfulBasis.LEGITIMATE_INTERESTS,
        consent_status=CandidateSource.ConsentStatus.NOT_REQUIRED,
        contact_permission=CandidateSource.ContactPermission.PERMITTED,
        recorded_by=user,
    )
    restricted = Candidate.objects.create(
        organization=organization,
        full_name="Application Only Candidate",
        created_by=user,
    )
    CandidateSource.objects.create(
        candidate=restricted,
        source_type=CandidateSource.SourceType.DOCUMENT_UPLOAD,
        source_name="Earlier application",
        lawful_basis=CandidateSource.LawfulBasis.LEGITIMATE_INTERESTS,
        consent_status=CandidateSource.ConsentStatus.NOT_REQUIRED,
        contact_permission=CandidateSource.ContactPermission.RESTRICTED,
        recorded_by=user,
    )
    inactive = Candidate.objects.create(
        organization=organization,
        full_name="Inactive Candidate",
        status=Candidate.Status.INACTIVE,
        created_by=user,
    )
    CandidateSource.objects.create(
        candidate=inactive,
        source_type=CandidateSource.SourceType.REFERRAL,
        source_name="Old referral",
        lawful_basis=CandidateSource.LawfulBasis.LEGITIMATE_INTERESTS,
        consent_status=CandidateSource.ConsentStatus.NOT_REQUIRED,
        contact_permission=CandidateSource.ContactPermission.PERMITTED,
        recorded_by=user,
    )
    already_added = Candidate.objects.create(
        organization=organization,
        full_name="Already Added Candidate",
        created_by=user,
    )
    associate_candidate_with_vacancy(
        candidate=already_added,
        vacancy=vacancy,
        user=user,
    )
    client.force_login(user)
    url = reverse(
        "candidates:vacancy-candidate-pool-add",
        args=[organization.slug, vacancy.pk],
    )

    page = client.get(url)
    content = page.content.decode()

    assert page.status_code == 200
    assert "Reusable Candidate" in content
    assert "Ready to add" in content
    assert "Application Only Candidate" in content
    assert "Application only" in content
    assert "Inactive Candidate" in content
    assert "Candidate is not active" in content
    assert "Already Added Candidate" not in content

    source_count = CandidateSource.objects.count()
    response = client.post(url, {"candidate_ids": [reusable.pk]})

    consideration = CandidateVacancyConsideration.objects.get(
        candidate=reusable,
        vacancy=vacancy,
    )
    assert response.status_code == 302
    assert response.url == reverse(
        "vacancies:vacancy-detail",
        args=[organization.slug, vacancy.pk],
    )
    assert consideration.origin == CandidateVacancyConsideration.Origin.CANDIDATE_POOL
    assert consideration.source == reusable_source
    assert consideration.created_by == user
    assert CandidateSource.objects.count() == source_count

    candidate_detail = client.get(
        reverse(
            "candidates:candidate-detail",
            args=[organization.slug, reusable.pk],
        )
    )
    assert "Added from candidate pool" in candidate_detail.content.decode()


def test_add_from_pool_rechecks_permission_and_tenant_boundaries(client):
    user, organization, vacancy = workspace_with_open_vacancy()
    restricted = Candidate.objects.create(
        organization=organization,
        full_name="Restricted Candidate",
        created_by=user,
    )
    CandidateSource.objects.create(
        candidate=restricted,
        source_type=CandidateSource.SourceType.DOCUMENT_UPLOAD,
        source_name="Application only",
        lawful_basis=CandidateSource.LawfulBasis.LEGITIMATE_INTERESTS,
        consent_status=CandidateSource.ConsentStatus.NOT_REQUIRED,
        contact_permission=CandidateSource.ContactPermission.RESTRICTED,
        recorded_by=user,
    )
    other_organization = Organization.objects.create(
        name="Other Organization",
        slug="other-organization",
    )
    other_candidate = Candidate.objects.create(
        organization=other_organization,
        full_name="Other Tenant Candidate",
    )
    client.force_login(user)
    url = reverse(
        "candidates:vacancy-candidate-pool-add",
        args=[organization.slug, vacancy.pk],
    )

    restricted_response = client.post(
        url,
        {"candidate_ids": [restricted.pk]},
        follow=True,
    )
    tenant_response = client.post(
        url,
        {"candidate_ids": [other_candidate.pk]},
        follow=True,
    )

    assert "Application only" in restricted_response.content.decode()
    assert "One or more selected candidates are unavailable" in (
        tenant_response.content.decode()
    )
    assert not CandidateVacancyConsideration.objects.filter(
        candidate__in=[restricted, other_candidate],
        vacancy=vacancy,
    ).exists()


def test_one_permitted_source_can_support_deliberate_reuse_for_two_vacancies():
    user, organization, first_vacancy = workspace_with_open_vacancy()
    second_vacancy = Vacancy.objects.create(
        organization=organization,
        title="Platform Engineer",
        description="Build platform services.",
        created_by=user,
    )
    requirements = VacancyRequirements.objects.create(
        vacancy=second_vacancy,
        source_description=second_vacancy.description,
        summary="Platform role",
        created_by=user,
    )
    confirm_requirements_and_open_vacancy(requirements=requirements, user=user)
    candidate = Candidate.objects.create(
        organization=organization,
        full_name="Reusable Across Roles",
        created_by=user,
    )
    source = CandidateSource.objects.create(
        candidate=candidate,
        source_type=CandidateSource.SourceType.REFERRAL,
        source_name="Synthetic referral",
        lawful_basis=CandidateSource.LawfulBasis.LEGITIMATE_INTERESTS,
        consent_status=CandidateSource.ConsentStatus.NOT_REQUIRED,
        contact_permission=CandidateSource.ContactPermission.PERMITTED,
        recorded_by=user,
    )

    first = add_candidates_from_pool(
        vacancy=first_vacancy,
        user=user,
        candidate_ids=[candidate.pk],
    )
    second = add_candidates_from_pool(
        vacancy=second_vacancy,
        user=user,
        candidate_ids=[candidate.pk],
    )

    assert first[0].source == source
    assert second[0].source == source
    assert candidate.vacancy_considerations.count() == 2
    assert candidate.sources.count() == 1


def test_vacancy_upload_inherits_policy_and_creates_scoped_application(
    client, settings, tmp_path
):
    settings.MEDIA_ROOT = tmp_path
    user, organization, vacancy = workspace_with_open_vacancy()
    OrganizationRetentionPolicy.objects.create(
        organization=organization,
        vacancy_candidate_lawful_basis=(
            OrganizationRetentionPolicy.CandidateLawfulBasis.LEGITIMATE_INTERESTS
        ),
    )
    client.force_login(user)

    response = client.post(
        intake_url(organization, vacancy),
        {"cv_files": [cv_upload()]},
    )
    batch = CandidateIntakeBatch.objects.get()
    item = batch.items.get()

    assert response.status_code == 302
    assert batch.vacancy == vacancy
    assert batch.source_name == "CV received for Senior Django Developer"
    assert batch.lawful_basis == CandidateSource.LawfulBasis.LEGITIMATE_INTERESTS
    assert batch.contact_permission == CandidateSource.ContactPermission.RESTRICTED
    assert batch.candidate_retention_until is None

    response = client.post(
        reverse(
            "candidates:candidate-intake-create-selected",
            args=[organization.slug, batch.pk],
        ),
        {
            f"item-{item.pk}-selected": "on",
            f"item-{item.pk}-full_name": "Candidate Person",
            f"item-{item.pk}-email": "candidate@example.test",
            f"item-{item.pk}-phone": "+383 44 123 456",
            f"item-{item.pk}-location": "Prishtina",
            f"item-{item.pk}-source_reference": "",
        },
    )

    candidate = Candidate.objects.get()
    consideration = CandidateVacancyConsideration.objects.get()
    item.refresh_from_db()
    assert response.status_code == 302
    assert consideration.candidate == candidate
    assert consideration.vacancy == vacancy
    assert consideration.document == item.accepted_document
    assert consideration.get_contact_scope_display() == "This vacancy only"
    assert assess_contact_permission(candidate=candidate, vacancy=vacancy).can_proceed
    assert not assess_contact_permission(candidate=candidate).can_proceed


def test_unambiguous_identity_match_reuses_candidate_for_new_vacancy_application(
    client, settings, tmp_path
):
    settings.MEDIA_ROOT = tmp_path
    user, organization, vacancy = workspace_with_open_vacancy()
    OrganizationRetentionPolicy.objects.create(
        organization=organization,
        vacancy_candidate_lawful_basis=(
            OrganizationRetentionPolicy.CandidateLawfulBasis.CONTRACT
        ),
    )
    existing = Candidate.objects.create(
        organization=organization,
        full_name="Existing Candidate",
        email="existing@example.test",
    )
    client.force_login(user)
    client.post(
        intake_url(organization, vacancy),
        {"cv_files": [cv_upload(name="new-version.docx", email=existing.email)]},
    )
    batch = CandidateIntakeBatch.objects.get()
    item = batch.items.get()

    response = client.post(
        reverse(
            "candidates:candidate-intake-create-selected",
            args=[organization.slug, batch.pk],
        ),
        {
            f"item-{item.pk}-selected": "on",
            f"item-{item.pk}-full_name": "Candidate Person",
            f"item-{item.pk}-email": existing.email,
            f"item-{item.pk}-phone": "+383 44 123 456",
            f"item-{item.pk}-location": "Prishtina",
            f"item-{item.pk}-source_reference": "",
        },
    )

    item.refresh_from_db()
    assert response.status_code == 302
    assert Candidate.objects.count() == 1
    assert item.candidate == existing
    assert item.accepted_document.candidate == existing
    assert CandidateVacancyConsideration.objects.get().candidate == existing


def test_exact_existing_cv_adds_application_without_storing_duplicate_bytes(
    settings, tmp_path
):
    settings.MEDIA_ROOT = tmp_path
    user, organization, vacancy = workspace_with_open_vacancy()
    existing = Candidate.objects.create(
        organization=organization,
        full_name="Existing Candidate",
    )
    original_upload = cv_upload()
    duplicate_upload = SimpleUploadedFile(
        original_upload.name,
        original_upload.read(),
        content_type=DOCX_CONTENT_TYPE,
    )
    original_upload.seek(0)
    document = upload_candidate_cv(
        candidate=existing,
        user=user,
        uploaded_file=original_upload,
    )
    batch = CandidateIntakeBatch.objects.create(
        organization=organization,
        vacancy=vacancy,
        source_name=f"CV received for {vacancy.title}",
        lawful_basis=CandidateSource.LawfulBasis.LEGITIMATE_INTERESTS,
        contact_permission=CandidateSource.ContactPermission.RESTRICTED,
        created_by=user,
    )

    item = upload_candidate_intake_cv(
        batch=batch,
        user=user,
        uploaded_file=duplicate_upload,
    )

    batch.refresh_from_db()
    consideration = CandidateVacancyConsideration.objects.get()
    assert Candidate.objects.count() == 1
    assert CandidateIntakeItem.objects.count() == 1
    assert item.status == CandidateIntakeItem.Status.CREATED
    assert item.candidate == existing
    assert item.accepted_document is None
    assert item.file.name == ""
    assert consideration.document == document
    assert batch.status == CandidateIntakeBatch.Status.COMPLETED
